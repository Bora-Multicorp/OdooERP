# -*- coding: utf-8 -*-
import base64
import re
from collections import defaultdict
from io import BytesIO

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class KsPoReceiptImportWizard(models.TransientModel):
    _name = "ks.po.receipt.import.wizard"
    _description = "Purchase Order Receipt Import Wizard (E-com)"

    file_data = fields.Binary(string="XLSX File")
    file_name = fields.Char(string="File Name")

    def action_download_template(self):
        return {
            "type": "ir.actions.act_url",
            "url": "/ks_ecom_purchase/static/src/xls/po_receipt_import_template.xlsx",
            "target": "self",
        }

    _ORDER_ID_ALIASES = {"orderid", "order_id", "orderno", "order_no", "ecomorderid"}
    _PRODUCT_CODE_ALIASES = {"asin", "sku", "productcode", "defaultcode", "internalreference", "barcode"}
    _QTY_DONE_ALIASES = {"orderquantity", "qtydone", "quantitydone", "receivedqty", "qtyreceived", "doneqty", "quantity"}

    @staticmethod
    def _normalize_header(value):
        return re.sub(r"[^a-z0-9]+", "", (value or "").strip().lower())

    @staticmethod
    def _to_string(value):
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _to_float(value, default=0.0):
        if value in (None, ""):
            return default
        try:
            return float(value)
        except Exception as exc:
            raise UserError(_("Invalid number value: %s") % value) from exc

    def _find_header(self, normalized_headers, aliases):
        for alias in aliases:
            if alias in normalized_headers:
                return alias
        return None

    def _extract_sheet_data(self):
        """Extract data from XLSX file"""
        self.ensure_one()
        if load_workbook is None:
            raise UserError(_("Python package 'openpyxl' is required to import XLSX files."))

        if not self.file_data or not self.file_name:
            raise UserError(_("Please upload an XLSX file before importing."))

        if not self.file_name.lower().endswith(".xlsx"):
            raise UserError(_("Please upload a valid .xlsx file only."))

        try:
            decoded_file = base64.b64decode(self.file_data)
        except Exception as exc:
            raise UserError(_("Unable to decode the uploaded file. Please upload it again.")) from exc

        try:
            workbook = load_workbook(filename=BytesIO(decoded_file), data_only=True)
            sheet = workbook.active
        except Exception as exc:
            raise UserError(_("Unable to read the XLSX file. Please verify file content.")) from exc

        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise UserError(_("The XLSX file does not contain importable rows."))

        headers = [self._to_string(cell) for cell in (rows[0] or [])]
        if not any(headers):
            raise UserError(_("Header row is empty in the uploaded file."))

        normalized_headers = {}
        for idx, header in enumerate(headers):
            normalized = self._normalize_header(header)
            if normalized:
                normalized_headers[normalized] = idx

        order_id_key = self._find_header(normalized_headers, self._ORDER_ID_ALIASES)
        asin_key = self._find_header(normalized_headers, self._PRODUCT_CODE_ALIASES)
        qty_done_key = self._find_header(normalized_headers, self._QTY_DONE_ALIASES)

        if not order_id_key:
            raise UserError(_("Missing required column: Order ID."))
        if not asin_key:
            raise UserError(_("Missing required column: ASIN/Product Code."))
        if not qty_done_key:
            raise UserError(_("Missing required column: Quantity Done/Received."))

        receipt_data = []
        for line_no, row in enumerate(rows[1:], start=2):
            if not row or all(cell in (None, "") for cell in row):
                continue

            def get_cell(key):
                idx = normalized_headers.get(key)
                if idx is None or idx >= len(row):
                    return None
                return row[idx]

            order_id = self._to_string(get_cell(order_id_key))
            asin = self._to_string(get_cell(asin_key))
            qty_done = self._to_float(get_cell(qty_done_key), default=0.0)

            if not order_id:
                raise UserError(_("Order ID is missing at row %s.") % line_no)
            if not asin:
                raise UserError(_("ASIN is missing at row %s.") % line_no)
            if qty_done <= 0:
                raise UserError(_("Quantity Done must be greater than 0 at row %s.") % line_no)

            # Capture all columns for E-com Information tab (field-value storage)
            row_raw = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                row_raw[header] = self._to_string(value)

            receipt_data.append({
                "line_no": line_no,
                "order_id": order_id,
                "asin": asin,
                "qty_done": qty_done,
                "raw": row_raw,
            })

        if not receipt_data:
            raise UserError(_("No valid import lines found in the uploaded file."))

        return receipt_data

    def _prepare_ecom_info_commands(self, sample_row):
        """Build commands to set ks_ecom_info_ids from a row's raw key-value data."""
        excluded = self._ORDER_ID_ALIASES | self._PRODUCT_CODE_ALIASES | self._QTY_DONE_ALIASES
        commands = []
        for key, value in (sample_row or {}).items():
            normalized = re.sub(r"[^a-z0-9]+", "", (key or "").strip().lower())
            if not normalized or normalized in excluded:
                continue
            if value in (None, ""):
                continue
            commands.append(
                Command.create({"field_key": key, "field_value": self._to_string(value)})
            )
        return commands

    def action_import_receipt_data(self):
        """Import receipt data: update existing stock pickings only (no new receipts created)."""
        self.ensure_one()
        receipt_data = self._extract_sheet_data()

        # Group by order_id for processing
        grouped_data = defaultdict(list)
        for row in receipt_data:
            grouped_data[row["order_id"]].append(row)

        updated_count = 0
        skipped_count = 0
        warnings = []

        for order_id, rows in grouped_data.items():
            # Find existing Purchase Order by unique identifier (ks_ecom_order_id or origin)
            po = self.env["purchase.order"].search([
                "|",
                ("ks_ecom_order_id", "=", order_id),
                ("origin", "ilike", order_id),
                ("state", "in", ("purchase", "done")),
            ], limit=1)
            # Prefer e-com imported PO when origin matches
            if not po or (po.origin and order_id in po.origin and not po.ks_ecom_order_id):
                po_ecom = self.env["purchase.order"].search([
                    ("ks_ecom_order_id", "=", order_id),
                    ("ks_ecom_order_id", "!=", False),
                ], limit=1)
                if po_ecom:
                    po = po_ecom

            if not po:
                po_by_name = self.env["purchase.order"].search([("name", "=", order_id)], limit=1)
                if po_by_name:
                    warnings.append(_("Order ID '%s' found (PO: %s) but not an e-com order. Only e-com imported orders can be updated.") % (order_id, po_by_name.name))
                else:
                    warnings.append(_("Order ID '%s' not found in Purchase Orders.") % order_id)
                skipped_count += len(rows)
                continue

            # Find existing receipt only (do not create new pickings)
            pickings = po.picking_ids.filtered(
                lambda p: p.picking_type_id.code == "incoming" and p.state in ("assigned", "confirmed", "draft", "done")
            )

            if not pickings:
                warnings.append(_("No existing receipt found for Order ID '%s' (PO: %s). Create receipt by confirming the PO first.") % (order_id, po.name))
                skipped_count += len(rows)
                continue

            # Use the first available picking (most recent)
            picking = pickings.sorted("create_date", reverse=True)[0]

            # Process each row for this order
            for row in rows:
                asin = row["asin"]
                qty_done = row["qty_done"]

                # Find product by ASIN (default_code)
                product = self.env["product.product"].search([
                    ("default_code", "=", asin)
                ], limit=1)

                if not product:
                    warnings.append(_("Product with ASIN '%s' not found (Order ID: %s, Row: %s).") % (
                        asin, order_id, row["line_no"]
                    ))
                    skipped_count += 1
                    continue

                # Find move in picking by product
                move = picking.move_ids.filtered(
                    lambda m: m.product_id.id == product.id and m.state != 'cancel'
                )

                if not move:
                    warnings.append(_("Product '%s' (ASIN: %s) not found in receipt for Order ID '%s' (Row: %s).") % (
                        product.name, asin, order_id, row["line_no"]
                    ))
                    skipped_count += 1
                    continue

                # Use the first move if multiple found
                move = move[0]

                # Find or create move line (stock.move.line) for this move
                move_line = move.move_line_ids.filtered(
                    lambda ml: ml.product_id.id == product.id
                )

                if not move_line:
                    move_line_vals = {
                        "move_id": move.id,
                        "product_id": product.id,
                        "product_uom_id": product.uom_id.id,
                        "location_id": move.location_id.id,
                        "location_dest_id": move.location_dest_id.id,
                        "quantity": qty_done,
                        "picked": True,
                    }
                    self.env["stock.move.line"].create(move_line_vals)
                    updated_count += 1
                else:
                    move_line.write({"quantity": qty_done, "picked": True})
                    updated_count += 1

            # Store Order ID and all imported Excel data in E-com Information (first row of this order)
            picking.ks_ecom_order_id = order_id
            info_commands = self._prepare_ecom_info_commands(rows[0].get("raw"))
            if info_commands:
                picking.ks_ecom_info_ids = [Command.clear()] + info_commands
            picking.is_ecom_updated = True

        # Same message style as PO Import: red box for records not updated, green box for success
        bus = self.env["bus.bus"]
        partner_id = self.env.user.partner_id

        # Red box: warnings / records not updated
        if warnings:
            warning_list = "; ".join(warnings[:10])
            if len(warnings) > 10:
                warning_list += _(" (and %s more)") % (len(warnings) - 10)
            bus._sendone(
                partner_id,
                "simple_notification",
                {
                    "type": "danger",
                    "title": _("Receipt Import: Records not updated"),
                    "message": warning_list,
                    "sticky": True,
                },
            )

        # Green box: success summary
        if updated_count > 0:
            bus._sendone(
                partner_id,
                "simple_notification",
                {
                    "type": "success",
                    "title": _("Receipt Import: Success"),
                    "message": _("Updated: %s line(s).") % updated_count,
                    "sticky": True,
                },
            )
        elif not warnings:
            bus._sendone(
                partner_id,
                "simple_notification",
                {
                    "type": "warning",
                    "title": _("Receipt Import"),
                    "message": _("No lines were updated. Check file format and data."),
                    "sticky": True,
                },
            )

        return {"type": "ir.actions.act_window_close"}

