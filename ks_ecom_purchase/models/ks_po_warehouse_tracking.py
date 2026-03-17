# -*- coding: utf-8 -*-

import base64
from io import BytesIO

from odoo import api, fields, models, _
from odoo.exceptions import UserError

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    Workbook = None


class KsPoWarehouseTracking(models.Model):
    _name = "ks.po.warehouse.tracking"
    _description = "PO Warehouse Tracking (E-com Import)"
    _order = "date desc, sr_no asc, id desc"

    sr_no = fields.Integer(
        string="SR No.",
        copy=False,
        readonly=True,
        help="Auto-incrementing sequence number.",
    )
    grn_id = fields.Many2one(
        "stock.picking",
        string="GRN No.",
        ondelete="set null",
        copy=False,
        help="Receipt (stock picking) linked to the PO.",
    )
    date = fields.Date(
        string="Date",
        copy=False,
        help="Purchase Order creation date.",
    )
    awb = fields.Char(string="AWB", default="111", copy=False)
    order_id = fields.Char(
        string="Order ID",
        copy=False,
        index=True,
        help="Value from PO's ks_ecom_order_id.",
    )
    invoice_no = fields.Char(
        string="Invoice No.",
        copy=False,
        help="Updated when the PO's Vendor Bill is posted.",
    )
    product_id = fields.Many2one(
        "product.product",
        string="Model",
        ondelete="restrict",
        copy=False,
    )
    qty = fields.Float(string="QTY", copy=False)
    rate = fields.Float(string="Rate", copy=False)
    iw = fields.Char(
        string="I/W",
        copy=False,
        help="Company name (from Excel 'Company' column, part before hyphen).",
    )
    person = fields.Char(string="Person", default="33", copy=False)
    delivered_qty = fields.Float(
        string="Delivered",
        copy=False,
        compute="_compute_delivered_qty_diff",
        store=True,
        help="Updated when the Receipt (GRN) is validated.",
    )
    diff = fields.Float(
        string="Diff",
        copy=False,
        compute="_compute_delivered_qty_diff",
        store=True,
        help="qty - delivered_qty.",
    )
    warehouse_id = fields.Many2one(
        "stock.warehouse",
        string="Warehouse",
        ondelete="restrict",
        index=True,
    )

    order_line_id = fields.Many2one(
        "purchase.order.line",
        string="PO Line",
        ondelete="cascade",
        index=True,
        copy=False,
        help="Link to purchase order line for invoice and delivery sync.",
    )

    @api.depends("order_line_id", "order_line_id.qty_received", "qty")
    def _compute_delivered_qty_diff(self):
        for rec in self:
            if rec.order_line_id:
                rec.delivered_qty = rec.order_line_id.qty_received
            else:
                rec.delivered_qty = 0.0
            rec.diff = rec.qty - rec.delivered_qty

    @api.model_create_multi
    def create(self, vals_list):
        seq = self.env["ir.sequence"].search(
            [("code", "=", "ks.po.warehouse.tracking")],
            limit=1,
        )
        for vals in vals_list:
            if vals.get("sr_no") is None and seq:
                try:
                    vals["sr_no"] = int(seq.next_by_id())
                except (TypeError, ValueError):
                    vals["sr_no"] = 0
        return super().create(vals_list)

    def action_export_xlsx(self):
        """Export selected or all tracking records to XLSX and return download action."""
        if Workbook is None:
            raise UserError(
                _("Python package 'openpyxl' is required to export XLSX. Please install it.")
            )
        records = self
        if not records:
            domain = self.env.context.get("search_domain", [])
            records = self.search(domain) if domain else self.search([])
        if not records:
            raise UserError(_("No records to export."))

        wb = Workbook()
        ws = wb.active
        ws.title = "PO Warehouse Tracking"

        headers = [
            "SR NO",
            "GRN NO",
            "DATE",
            "AWB",
            "Order ID",
            "Invoice No",
            "Model",
            "QTY",
            "Rate",
            "I/W",
            "Person",
            "Delivered",
            "Diff",
            "Warehouse",
        ]
        header_font = Font(bold=True)
        for col, label in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = header_font

        for row_idx, rec in enumerate(records, start=2):
            ws.cell(row=row_idx, column=1, value=rec.sr_no or "")
            ws.cell(row=row_idx, column=2, value=rec.grn_id.name if rec.grn_id else "")
            ws.cell(row=row_idx, column=3, value=rec.date.isoformat() if rec.date else "")
            ws.cell(row=row_idx, column=4, value=rec.awb or "")
            ws.cell(row=row_idx, column=5, value=rec.order_id or "")
            ws.cell(row=row_idx, column=6, value=rec.invoice_no or "")
            ws.cell(row=row_idx, column=7, value=rec.product_id.display_name if rec.product_id else "")
            ws.cell(row=row_idx, column=8, value=rec.qty)
            ws.cell(row=row_idx, column=9, value=rec.rate)
            ws.cell(row=row_idx, column=10, value=rec.iw or "")
            ws.cell(row=row_idx, column=11, value=rec.person or "")
            ws.cell(row=row_idx, column=12, value=rec.delivered_qty)
            ws.cell(row=row_idx, column=13, value=rec.diff)
            ws.cell(row=row_idx, column=14, value=rec.warehouse_id.name if rec.warehouse_id else "")

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        data = base64.b64encode(buffer.getvalue())

        name = "po_warehouse_tracking_export.xlsx"
        attachment = self.env["ir.attachment"].create({
            "name": name,
            "datas": data,
            "res_model": self._name,
            "res_id": 0,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })
        return {
            "type": "ir.actions.act_url",
            "url": "/web/content/%s?download=1" % attachment.id,
            "target": "self",
        }
