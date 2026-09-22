# -*- coding: utf-8 -*-
import base64
import io
import re

from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    openpyxl = None


class StockPickingUploadExcelWizard(models.TransientModel):
    _name = "stock.picking.upload.excel.wizard"
    _description = "Upload Serials/Lots and IMEIs from Excel for Stock Picking"

    picking_id = fields.Many2one(
        "stock.picking",
        string="Transfer / Picking",
        required=True,
        ondelete="cascade",
        default=lambda self: self.env.context.get("active_id") or self.env.context.get("default_picking_id"),
    )
    picking_type_code = fields.Selection(
        related="picking_id.picking_type_id.code",
        string="Operation Type",
        readonly=True,
    )
    excel_file = fields.Binary(string="Excel or CSV File", required=False)
    filename = fields.Char(string="Filename")
    keep_lines = fields.Boolean(
        string="Keep existing lines",
        default=True,
        help="1.checked, existing lines with serial numbers/IMEIs are kept and newly uploaded lines are added. 2. unchecked, existing lines for the moves in the file will be removed and replaced by the uploaded lines.",
    )
    test_passed = fields.Boolean(
        string="File Tested",
        default=False,
    )
    test_message = fields.Char(
        string="Test Result",
        readonly=True,
    )

    SUPPORTED_EXTENSIONS = (".xlsx", ".xls", ".csv")

    @api.model
    def default_get(self, fields_list):
        res = super().default_get(fields_list)
        res["keep_lines"] = True
        res["test_passed"] = False
        res["test_message"] = False
        return res

    @api.onchange("excel_file")
    def _onchange_excel_file(self):
        self.test_passed = False
        self.test_message = False

    # -------------------------------------------------------------------------
    # SAMPLE FILE DOWNLOAD
    # -------------------------------------------------------------------------
    def action_download_sample_xlsx(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("Excel export requires the 'openpyxl' Python library."))
        if not self.picking_id:
            raise UserError(_("No transfer / picking specified."))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Serials"

        # Reference sheet for Country Dropdown
        ws_countries = wb.create_sheet(title="Countries")
        countries = self.env["res.country"].search([("name", "!=", False)], order="name asc")
        country_count = 0
        for idx, country in enumerate(countries, start=1):
            ws_countries.cell(row=idx, column=1, value=country.name)
            country_count += 1

        # Hide the Countries reference sheet so only the main template is displayed to users
        ws_countries.sheet_state = "hidden"

        picking = self.picking_id
        is_outgoing = picking.picking_type_id.code == "outgoing"
        picking_ext_id = self._get_or_create_external_id(picking)

        # Fallback country from picking header if set
        picking_country_name = picking.made_country.name if picking.made_country else ""

        current_row = 2
        moves = picking.move_ids.filtered(lambda m: m.state != "cancel")

        # Determine remaining quantity for each move (demand minus already assigned serial/lot/IMEI lines)
        move_remaining = {}
        for move in moves:
            product = move.product_id
            assigned_lines = move.move_line_ids.filtered(
                lambda l: bool(l.lot_name or l.lot_id or l.imei or (l.quantity and l.quantity > 0))
            )
            assigned_qty = int(sum(assigned_lines.mapped("quantity")) or len(assigned_lines))
            already_assigned_count = max(len(assigned_lines), assigned_qty)
            if product.tracking in ("serial", "lot"):
                demand_qty = int(move.product_uom_qty or 0)
                if demand_qty <= 0:
                    demand_qty = int(move.quantity or 0)
                rem = max(0, demand_qty - already_assigned_count)
                if demand_qty == 0 and already_assigned_count == 0:
                    rem = 1
            else:
                demand_qty = int(move.product_uom_qty or move.quantity or 1)
                rem = max(0, demand_qty - already_assigned_count)
            move_remaining[move] = (already_assigned_count, rem)

        # IMEI 2 column is included IF AND ONLY IF any product with REMAINING lines has dual SIM enabled
        has_dual_sim = any(
            self._is_product_dual_sim(m.product_id)
            for m in moves
            if move_remaining.get(m, (0, 0))[1] > 0
        )
        if not any(rem > 0 for _, rem in move_remaining.values()):
            has_dual_sim = any(self._is_product_dual_sim(m.product_id) for m in moves)

        if has_dual_sim:
            headers = [
                "External ID : Receipts",
                "External ID : Move ID",
                "Product Name",
                "Serial Number",
                "IMEI 1",
                "IMEI 2",
                "Made In",
            ]
        else:
            headers = [
                "External ID : Receipts",
                "External ID : Move ID",
                "Product Name",
                "Serial Number",
                "IMEI 1",
                "Made In",
            ]

        # Setup styles
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_side = Side(border_style="thin", color="D3D3D3")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.row_dimensions[1].height = 28
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        for move in moves:
            already_assigned_count, remaining_qty = move_remaining[move]
            if remaining_qty <= 0:
                continue

            product = move.product_id
            is_mobile = self._is_product_mobile(product)
            is_dual = self._is_product_dual_sim(product)
            move_ext_id = self._get_or_create_external_id(move)

            # Made in country pre-fill (from move or picking)
            row_country_name = (
                (move.made_in_country_id and move.made_in_country_id.name)
                or (move.made_country and move.made_country.name)
                or picking_country_name
                or ""
            )

            # For outgoing, find quants/lots in stock excluding already assigned lots
            available_quants = []
            if is_outgoing:
                assigned_lines = move.move_line_ids.filtered(
                    lambda l: bool(l.lot_name or l.lot_id or l.imei)
                )
                assigned_lot_ids = assigned_lines.mapped("lot_id").ids
                quant_domain = [
                    ("product_id", "=", product.id),
                    ("location_id", "child_of", move.location_id.id),
                    ("quantity", ">", 0),
                    ("lot_id", "!=", False),
                ]
                if assigned_lot_ids:
                    quant_domain.append(("lot_id", "not in", assigned_lot_ids))
                available_quants = self.env["stock.quant"].search(quant_domain, limit=remaining_qty)

            for offset in range(remaining_qty):
                seq = already_assigned_count + offset + 1
                # 1. Receipt External ID
                receipt_val = picking_ext_id

                # 2. Move External ID
                move_val = move_ext_id

                # 3. Product Name
                product_val = product.display_name or product.name

                # 4. Serial Number
                if is_outgoing and offset < len(available_quants):
                    q = available_quants[offset]
                    serial_val = q.lot_id.name
                    imei1_val = q.imei or ""
                    imei2_val = q.imei2 or ""
                else:
                    serial_val = f"SN{move.id}{seq:04d}"
                    imei1_val = ""
                    imei2_val = ""

                # 5 & 6. IMEI 1 and IMEI 2 based on conditions
                if is_mobile:
                    if not imei1_val:
                        imei1_val = f"86{move.id % 10000:04d}{seq:09d}"
                    if is_dual and not imei2_val:
                        imei2_val = f"87{move.id % 10000:04d}{seq:09d}"
                else:
                    imei1_val = ""
                    imei2_val = ""

                if has_dual_sim:
                    row_vals = [
                        receipt_val,
                        move_val,
                        product_val,
                        serial_val,
                        imei1_val,
                        imei2_val,
                        row_country_name,
                    ]
                    text_col_indices = (4, 5, 6)
                    center_col_indices = (1, 2, 4, 5, 6, 7)
                else:
                    row_vals = [
                        receipt_val,
                        move_val,
                        product_val,
                        serial_val,
                        imei1_val,
                        row_country_name,
                    ]
                    text_col_indices = (4, 5)
                    center_col_indices = (1, 2, 4, 5, 6)

                ws.row_dimensions[current_row].height = 20
                for c_idx, val in enumerate(row_vals, 1):
                    c = ws.cell(row=current_row, column=c_idx, value=val)
                    c.border = border
                    # Force IMEI and Serial as text format so Excel does not convert to scientific notation
                    if c_idx in text_col_indices:
                        c.number_format = "@"
                    if c_idx in center_col_indices:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

                current_row += 1

        # Fallback if picking had no moves or all were completed
        if current_row == 2:
            first_move = moves[0] if moves else False
            first_prod = first_move.product_id if first_move else False
            if has_dual_sim:
                sample_row = [
                    picking_ext_id or "WH/IN/00001",
                    self._get_or_create_external_id(first_move) if first_move else "stock_move_1",
                    (first_prod and (first_prod.display_name or first_prod.name)) or "Sample Mobile Product",
                    "SN0000001",
                    "864201040000001",
                    "864201040000002",
                    picking_country_name,
                ]
                sample_text_cols = (4, 5, 6)
            else:
                sample_row = [
                    picking_ext_id or "WH/IN/00001",
                    self._get_or_create_external_id(first_move) if first_move else "stock_move_1",
                    (first_prod and (first_prod.display_name or first_prod.name)) or "Sample Mobile Product",
                    "SN0000001",
                    "864201040000001",
                    picking_country_name,
                ]
                sample_text_cols = (4, 5)

            for c_idx, val in enumerate(sample_row, 1):
                c = ws.cell(row=2, column=c_idx, value=val)
                c.border = border
                if c_idx in sample_text_cols:
                    c.number_format = "@"
            current_row = 3

        max_validation_row = max(current_row + 100, 500)

        # Add Data Validation for IMEI 1 (Column E) - must be exactly 15 digits
        dv_imei1 = DataValidation(
            type="textLength",
            operator="equal",
            formula1="15",
            allow_blank=True,
            showErrorMessage=True,
            showInputMessage=True,
            errorTitle=_("Invalid IMEI 1"),
            error=_("IMEI 1 must be exactly 15 digits."),
            promptTitle=_("IMEI 1"),
            prompt=_("Enter a 15-digit IMEI number."),
        )
        ws.add_data_validation(dv_imei1)
        dv_imei1.add(f"E2:E{max_validation_row}")

        # Add Data Validation for IMEI 2 (Column F) if dual SIM is present - must be exactly 15 digits
        if has_dual_sim:
            dv_imei2 = DataValidation(
                type="textLength",
                operator="equal",
                formula1="15",
                allow_blank=True,
                showErrorMessage=True,
                showInputMessage=True,
                errorTitle=_("Invalid IMEI 2"),
                error=_("IMEI 2 must be exactly 15 digits."),
                promptTitle=_("IMEI 2"),
                prompt=_("Enter a 15-digit IMEI number."),
            )
            ws.add_data_validation(dv_imei2)
            dv_imei2.add(f"F2:F{max_validation_row}")

        # Add Data Validation Dropdown for "Made In" (last column)
        if country_count > 0:
            made_in_col_letter = get_column_letter(len(headers))
            dv = DataValidation(
                type="list",
                formula1=f"=Countries!$A$1:$A${country_count}",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle=_("Invalid Country"),
                error=_("Please select a country from the dropdown list."),
            )
            ws.add_data_validation(dv)
            dv.add(f"{made_in_col_letter}2:{made_in_col_letter}{max_validation_row}")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        xlsx_data = output.read()

        clean_picking_name = re.sub(r"[^\w\-]", "_", picking.name or "picking")
        filename = f"sample_serials_{clean_picking_name}.xlsx"

        attachment = self.env["ir.attachment"].sudo().create({
            "name": filename,
            "type": "binary",
            "datas": base64.b64encode(xlsx_data),
            "res_model": "stock.picking",
            "res_id": picking.id,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "public": True,
        })

        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}/{filename}?download=true",
            "target": "new",
        }

    # -------------------------------------------------------------------------
    # IMPORT PROCESSING
    # -------------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        if not self.picking_id:
            raise UserError(_("No transfer / picking specified."))
        if not self.excel_file or not self.filename:
            raise ValidationError(_("Please select and upload a file (.xlsx, .xls, or .csv)."))

        if not self.test_passed:
            self.action_test_file()

        fn_lower = self.filename.lower()
        if not any(fn_lower.endswith(ext) for ext in self.SUPPORTED_EXTENSIONS):
            raise ValidationError(_("Unsupported file format. Please upload .xlsx, .xls or .csv file."))

        try:
            file_content = base64.b64decode(self.excel_file)
        except Exception as e:
            raise ValidationError(_("Could not decode file: %s") % e)

        if fn_lower.endswith(".csv"):
            parsed_rows = self._parse_csv_file(file_content)
        else:
            parsed_rows = self._parse_xlsx_file(file_content)

        if not parsed_rows:
            raise ValidationError(_("The uploaded file contains no data rows."))

        # Extract structured row dicts
        structured_rows = self._extract_structured_rows(parsed_rows)
        if not structured_rows:
            raise ValidationError(_("No valid data rows found in the uploaded file."))

        # Apply serials, IMEIs, and Made In to the picking moves
        self._apply_rows_to_picking(structured_rows)

        return {"type": "ir.actions.act_window_close"}

    def action_test_file(self):
        """Verify that the selected Excel/CSV file matches the current transfer's External ID and Move IDs."""
        self.ensure_one()
        if not self.excel_file:
            raise ValidationError(_("Please select and upload a file (.xlsx, .xls, or .csv) to test."))

        fn_lower = (self.filename or "").lower()
        if not any(fn_lower.endswith(ext) for ext in self.SUPPORTED_EXTENSIONS):
            raise ValidationError(_("Unsupported file format. Please upload .xlsx, .xls or .csv file."))

        try:
            file_content = base64.b64decode(self.excel_file)
        except Exception as e:
            raise ValidationError(_("Could not decode file: %s") % e)

        if fn_lower.endswith(".csv"):
            parsed_rows = self._parse_csv_file(file_content)
        else:
            parsed_rows = self._parse_xlsx_file(file_content)

        if not parsed_rows:
            raise ValidationError(_("The selected file contains no data rows."))

        structured_rows = self._extract_structured_rows(parsed_rows)
        if not structured_rows:
            raise ValidationError(_("No valid data rows found in the selected file."))

        picking = self.picking_id
        picking_ext_id = self._get_or_create_external_id(picking)
        current_move_ids = set(picking.move_ids.ids)
        current_move_ext_ids = {self._get_or_create_external_id(m) for m in picking.move_ids}

        errors = []
        matching_rows = 0

        for r in structured_rows:
            row_num = r["row_number"]
            receipt_ref = r["receipt_ref"]
            move_ref = r["move_ref"]
            product_ref = r["product_ref"]
            imei1 = r["imei1"]
            imei2 = r["imei2"]

            # 1. Check Receipt External ID match
            if receipt_ref:
                resolved_picking = self._resolve_record_by_ref("stock.picking", receipt_ref)
                if resolved_picking and resolved_picking.id != picking.id:
                    errors.append(
                        _("Row %s: External ID : Receipts '%s' belongs to transfer '%s', NOT current transfer '%s'.")
                        % (row_num, receipt_ref, resolved_picking.name, picking.name)
                    )
                elif not resolved_picking and receipt_ref not in (picking_ext_id, picking.name):
                    errors.append(
                        _("Row %s: External ID : Receipts '%s' does not match current transfer '%s'.")
                        % (row_num, receipt_ref, picking.name)
                    )

            # 2. Check Move External ID / Product match
            move = None
            if move_ref:
                resolved_move = self._resolve_record_by_ref("stock.move", move_ref)
                if resolved_move and resolved_move.id in current_move_ids:
                    move = resolved_move
                elif resolved_move and resolved_move.id not in current_move_ids:
                    errors.append(
                        _("Row %s: Move ID '%s' belongs to transfer '%s', not current transfer '%s'.")
                        % (row_num, move_ref, resolved_move.picking_id.name, picking.name)
                    )
                elif move_ref not in current_move_ext_ids:
                    errors.append(
                        _("Row %s: Move ID '%s' not found in current transfer '%s'.")
                        % (row_num, move_ref, picking.name)
                    )

            if not move and product_ref:
                matching_moves = picking.move_ids.filtered(
                    lambda m: m.product_id.name == product_ref
                    or m.product_id.display_name == product_ref
                    or m.product_id.default_code == product_ref
                )
                if not matching_moves and not move_ref:
                    errors.append(
                        _("Row %s: Product '%s' not found in current transfer '%s'.")
                        % (row_num, product_ref, picking.name)
                    )

            # 3. Check IMEI length constraints
            if imei1 and (not imei1.isdigit() or len(imei1) != 15):
                errors.append(
                    _("Row %s: IMEI 1 '%s' must be a 15-digit number.") % (row_num, imei1)
                )
            if imei2 and (not imei2.isdigit() or len(imei2) != 15):
                errors.append(
                    _("Row %s: IMEI 2 '%s' must be a 15-digit number.") % (row_num, imei2)
                )

            matching_rows += 1

        if errors:
            self.write({
                "test_passed": False,
                "test_message": False,
            })
            err_msg = _("Test File Failed - Found %s issue(s):\n\n") % len(errors)
            err_msg += "\n".join(errors[:10])
            if len(errors) > 10:
                err_msg += _("\n... and %s more issue(s).") % (len(errors) - 10)
            raise ValidationError(err_msg)

        msg = _("Receipt External ID matches '%s'. All %s data row(s) are valid and ready to upload.") % (
            picking.name, matching_rows
        )
        self.write({
            "test_passed": True,
            "test_message": msg,
        })

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "views": [[False, "form"]],
            "target": "new",
            "context": self.env.context,
        }

    def _parse_xlsx_file(self, file_content):
        if openpyxl is None:
            raise ValidationError(_("Excel (.xlsx) support requires the 'openpyxl' library."))
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
        except Exception as e:
            raise ValidationError(_("Could not open Excel file: %s") % e)

        rows = []
        try:
            ws = wb.active
            if not ws:
                return []
            for row in ws.iter_rows(min_row=1, values_only=True):
                if not row:
                    continue
                row_str_vals = [self._cell_to_str(c) for c in row]
                if any(row_str_vals):
                    rows.append(row_str_vals)
        finally:
            wb.close()
        return rows

    def _parse_csv_file(self, file_content):
        try:
            text = file_content.decode("utf-8", errors="replace")
        except Exception as e:
            raise ValidationError(_("Could not read CSV file: %s") % e)

        lines = [line.strip() for line in text.splitlines() if line.strip()]
        rows = []
        for line in lines:
            parts = self._parse_csv_line(line)
            if any(parts):
                rows.append([self._cell_to_str(p) for p in parts])
        return rows

    def _parse_csv_line(self, line):
        result = []
        current = []
        in_quotes = False
        for c in line:
            if c == '"':
                in_quotes = not in_quotes
            elif not in_quotes and c in (",", "\t"):
                result.append("".join(current).strip())
                current = []
            else:
                current.append(c)
        result.append("".join(current).strip())
        return result

    def _cell_to_str(self, value):
        if value is None:
            return ""
        if isinstance(value, bool):
            return "True" if value else "False"
        if isinstance(value, float):
            return str(int(value)) if value.is_integer() else str(value).strip()
        if isinstance(value, int):
            return str(value).strip()
        return str(value).strip()

    def _is_product_mobile(self, product):
        if not product:
            return False
        tmpl = product.product_tmpl_id or product
        return bool(
            getattr(product, "is_mobile_category_selected", False)
            or getattr(tmpl, "is_mobile_category_selected", False)
        )

    def _is_product_dual_sim(self, product):
        if not product:
            return False
        if not self._is_product_mobile(product):
            return False
        tmpl = product.product_tmpl_id or product
        return bool(
            getattr(product, "is_dual_sim", False)
            or getattr(tmpl, "is_dual_sim", False)
        )

    def _extract_structured_rows(self, parsed_rows):
        first_row = parsed_rows[0]
        has_header = self._is_header_row(first_row)

        if has_header:
            header_map = self._map_header_indices(first_row)
            data_rows = parsed_rows[1:]
        else:
            if len(first_row) >= 7:
                header_map = {
                    "receipt_idx": 0,
                    "move_idx": 1,
                    "product_idx": 2,
                    "serial_idx": 3,
                    "imei1_idx": 4,
                    "imei2_idx": 5,
                    "made_in_idx": 6,
                }
            else:
                header_map = {
                    "receipt_idx": 0,
                    "move_idx": 1,
                    "product_idx": 2,
                    "serial_idx": 3,
                    "imei1_idx": 4,
                    "imei2_idx": -1,
                    "made_in_idx": 5,
                }
            data_rows = parsed_rows

        structured = []
        for row_idx, row in enumerate(data_rows, start=2 if has_header else 1):
            def get_col(idx_key):
                idx = header_map.get(idx_key, -1)
                if idx != -1 and idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
                return ""

            receipt_val = get_col("receipt_idx")
            move_val = get_col("move_idx")
            product_val = get_col("product_idx")
            serial_val = get_col("serial_idx")
            imei1_val = get_col("imei1_idx")
            imei2_val = get_col("imei2_idx")
            made_in_val = get_col("made_in_idx")

            # Ignore completely empty rows
            if not any([receipt_val, move_val, product_val, serial_val, imei1_val, imei2_val, made_in_val]):
                continue

            structured.append({
                "row_number": row_idx,
                "receipt_ref": receipt_val,
                "move_ref": move_val,
                "product_ref": product_val,
                "serial_number": serial_val,
                "imei1": imei1_val,
                "imei2": imei2_val,
                "made_in": made_in_val,
            })
        return structured

    def _is_header_row(self, row):
        if not row:
            return False
        row_str = " ".join(str(c or "").lower() for c in row)
        return any(k in row_str for k in ("receipt", "move", "serial", "lot", "imei", "product", "made in"))

    def _map_header_indices(self, headers):
        mapping = {
            "receipt_idx": -1,
            "move_idx": -1,
            "product_idx": -1,
            "serial_idx": -1,
            "imei1_idx": -1,
            "imei2_idx": -1,
            "made_in_idx": -1,
        }

        for idx, h in enumerate(headers):
            h_clean = str(h or "").strip().lower()
            if not h_clean:
                continue

            if ("receipt" in h_clean or "external id : receipts" in h_clean or "external id for the receipts" in h_clean or "picking" in h_clean) and mapping["receipt_idx"] == -1:
                mapping["receipt_idx"] = idx
            elif ("move" in h_clean or "external id : move id" in h_clean or "external id: move id" in h_clean) and mapping["move_idx"] == -1:
                mapping["move_idx"] = idx
            elif "product" in h_clean and mapping["product_idx"] == -1:
                mapping["product_idx"] = idx
            elif ("imei 2" in h_clean or "imei2" in h_clean or "imei_2" in h_clean) and mapping["imei2_idx"] == -1:
                mapping["imei2_idx"] = idx
            elif ("imei 1" in h_clean or "imei1" in h_clean or "imei_1" in h_clean or h_clean == "imei") and mapping["imei1_idx"] == -1:
                mapping["imei1_idx"] = idx
            elif ("serial" in h_clean or "lot" in h_clean) and mapping["serial_idx"] == -1:
                mapping["serial_idx"] = idx
            elif ("made" in h_clean or "country" in h_clean or "origin" in h_clean) and mapping["made_in_idx"] == -1:
                mapping["made_in_idx"] = idx

        # Set of indices already claimed
        used_indices = {v for v in mapping.values() if v != -1}

        # If made_in_idx was not matched by keyword, check if the last column is available
        last_col_idx = len(headers) - 1
        if mapping["made_in_idx"] == -1 and last_col_idx not in used_indices:
            mapping["made_in_idx"] = last_col_idx
            used_indices.add(last_col_idx)

        # For remaining unmapped fields, fill ONLY from unused column indices
        available_indices = [i for i in range(len(headers)) if i not in used_indices]
        order_to_fill = ["receipt_idx", "move_idx", "product_idx", "serial_idx", "imei1_idx"]
        if len(headers) >= 7:
            order_to_fill.append("imei2_idx")
        order_to_fill.append("made_in_idx")

        for k in order_to_fill:
            if mapping[k] == -1 and available_indices:
                mapping[k] = available_indices.pop(0)

        return mapping

    def _find_country(self, country_str):
        if not country_str:
            return self.env["res.country"]
        country_str = str(country_str).strip()
        Country = self.env["res.country"]
        c = Country.search([("name", "=ilike", country_str)], limit=1)
        if c:
            return c
        c = Country.search([("code", "=ilike", country_str)], limit=1)
        if c:
            return c
        if country_str.isdigit():
            c = Country.browse(int(country_str))
            if c.exists():
                return c
        return Country

    def _get_or_create_external_id(self, record):
        """Return the External ID (XML ID) for record, creating an __export__ one if none exists."""
        if not record:
            return ""
        ext_ids = record.get_external_id()
        if ext_ids.get(record.id):
            return ext_ids[record.id]

        imd = self.env["ir.model.data"].sudo().search([
            ("model", "=", record._name),
            ("res_id", "=", record.id),
        ], limit=1)
        if imd:
            return f"{imd.module}.{imd.name}" if imd.module else imd.name

        table_name = record._table if hasattr(record, "_table") else record._name.replace(".", "_")
        xml_name = f"{table_name}_{record.id}"
        existing = self.env["ir.model.data"].sudo().search([
            ("module", "=", "__export__"),
            ("name", "=", xml_name),
        ], limit=1)
        if existing and existing.res_id != record.id:
            xml_name = f"{table_name}_{record.id}_{record.id}"

        self.env["ir.model.data"].sudo().create({
            "name": xml_name,
            "module": "__export__",
            "model": record._name,
            "res_id": record.id,
        })
        return f"__export__.{xml_name}"

    def _resolve_record_by_ref(self, model_name, ref_str):
        """Find a record by External ID (XML ID) or integer ID."""
        if not ref_str:
            return self.env[model_name]
        ref_str = str(ref_str).strip()

        # 1. External ID with module prefix (e.g. __export__.stock_move_123 or base.something)
        if "." in ref_str:
            try:
                rec = self.env.ref(ref_str, raise_if_not_found=False)
                if rec and rec._name == model_name and rec.exists():
                    return rec
            except Exception:
                pass
            parts = ref_str.split(".", 1)
            imd = self.env["ir.model.data"].sudo().search([
                ("module", "=", parts[0]),
                ("name", "=", parts[1]),
                ("model", "=", model_name),
            ], limit=1)
            if imd and imd.res_id:
                rec = self.env[model_name].browse(imd.res_id)
                if rec.exists():
                    return rec

        # 2. Search by XML ID name without module prefix
        imd = self.env["ir.model.data"].sudo().search([
            ("name", "=", ref_str),
            ("model", "=", model_name),
        ], limit=1)
        if imd and imd.res_id:
            rec = self.env[model_name].browse(imd.res_id)
            if rec.exists():
                return rec

        # 3. Integer database ID
        if ref_str.isdigit():
            rec = self.env[model_name].browse(int(ref_str))
            if rec.exists():
                return rec

        return self.env[model_name]

    def _apply_rows_to_picking(self, rows):
        picking = self.picking_id
        is_outgoing = picking.picking_type_id.code == "outgoing"

        # 1. Match moves & validate rows
        seen_serials = set()
        seen_imei1 = set()
        seen_imei2 = set()

        # Existing serials and IMEIs on this transfer (preserve manual entries)
        existing_serials = set()
        existing_imeis = set()
        if self.keep_lines:
            for m in picking.move_ids:
                for l in m.move_line_ids:
                    s = (l.lot_name or (l.lot_id and l.lot_id.name) or "").strip()
                    if s:
                        existing_serials.add(s)
                    if l.imei:
                        existing_imeis.add(str(l.imei).strip())
                    if l.imei2:
                        existing_imeis.add(str(l.imei2).strip())

        rows_by_move = {}
        for r in rows:
            row_num = r["row_number"]
            move_ref = r["move_ref"]
            product_ref = r["product_ref"]
            serial = r["serial_number"]
            imei1 = r["imei1"]
            imei2 = r["imei2"]
            made_in = r["made_in"]

            # Locate move
            move = None
            if move_ref:
                resolved_move = self._resolve_record_by_ref("stock.move", move_ref)
                if resolved_move and resolved_move.id in picking.move_ids.ids:
                    move = resolved_move

            if not move and product_ref:
                # Fallback to product display name or name in picking moves
                matching_moves = picking.move_ids.filtered(
                    lambda m: m.product_id.name == product_ref
                    or m.product_id.display_name == product_ref
                    or m.product_id.default_code == product_ref
                )
                if matching_moves:
                    move = matching_moves[0]

            if not move:
                raise ValidationError(
                    _("Row %s: Could not find matching Stock Move in transfer %s for Move ID '%s' / Product '%s'.")
                    % (row_num, picking.name, move_ref, product_ref)
                )

            product = move.product_id
            is_mobile = self._is_product_mobile(product)
            is_dual = self._is_product_dual_sim(product)

            # Serial number check
            if product.tracking in ("serial", "lot") and not serial:
                raise ValidationError(
                    _("Row %s: Serial Number is required for tracked product '%s'.")
                    % (row_num, product.display_name)
                )

            # Duplicate serial in uploaded file or already assigned on picking
            if serial:
                if serial in seen_serials:
                    raise ValidationError(_("Row %s: Serial Number '%s' is duplicated in the file.") % (row_num, serial))
                if serial in existing_serials:
                    raise ValidationError(_("Row %s: Serial Number '%s' is already assigned in this transfer.") % (row_num, serial))
                seen_serials.add(serial)

            # IMEI 1 validation
            if is_mobile:
                if not imei1:
                    raise ValidationError(
                        _("Row %s: IMEI 1 is mandatory for mobile product '%s'.") % (row_num, product.display_name)
                    )
                if not imei1.isdigit() or len(imei1) != 15:
                    raise ValidationError(
                        _("Row %s: IMEI 1 must be a 15-digit number. Got: '%s'") % (row_num, imei1)
                    )
                if imei1 in seen_imei1 or imei1 in seen_imei2:
                    raise ValidationError(_("Row %s: IMEI 1 '%s' is duplicated in the file.") % (row_num, imei1))
                if imei1 in existing_imeis:
                    raise ValidationError(_("Row %s: IMEI 1 '%s' is already assigned in this transfer.") % (row_num, imei1))
                seen_imei1.add(imei1)
            elif imei1:
                # If non-mobile but IMEI 1 entered, must be 15 digits
                if not imei1.isdigit() or len(imei1) != 15:
                    raise ValidationError(
                        _("Row %s: IMEI 1 must be a 15-digit number. Got: '%s'") % (row_num, imei1)
                    )
                if imei1 in existing_imeis:
                    raise ValidationError(_("Row %s: IMEI 1 '%s' is already assigned in this transfer.") % (row_num, imei1))

            # IMEI 2 validation
            if is_mobile and is_dual:
                if not imei2:
                    raise ValidationError(
                        _("Row %s: IMEI 2 is mandatory for dual-SIM mobile product '%s'.")
                        % (row_num, product.display_name)
                    )
                if not imei2.isdigit() or len(imei2) != 15:
                    raise ValidationError(
                        _("Row %s: IMEI 2 must be a 15-digit number. Got: '%s'") % (row_num, imei2)
                    )
                if imei2 in seen_imei1 or imei2 in seen_imei2:
                    raise ValidationError(_("Row %s: IMEI 2 '%s' is duplicated in the file.") % (row_num, imei2))
                if imei2 in existing_imeis:
                    raise ValidationError(_("Row %s: IMEI 2 '%s' is already assigned in this transfer.") % (row_num, imei2))
                seen_imei2.add(imei2)

                # IMEI 1 and 2 must differ (unless Samsung / OnePlus)
                if imei1 and imei2 and imei1 == imei2:
                    brand_name = (product.product_tmpl_id.brand_id.name or "").lower()
                    if brand_name not in ("samsung", "oneplus"):
                        raise ValidationError(_("Row %s: IMEI 1 and IMEI 2 must be different.") % row_num)
            elif imei2:
                # Non-dual or non-mobile with IMEI 2 entered
                if not imei2.isdigit() or len(imei2) != 15:
                    raise ValidationError(
                        _("Row %s: IMEI 2 must be a 15-digit number. Got: '%s'") % (row_num, imei2)
                    )
                if imei2 in existing_imeis:
                    raise ValidationError(_("Row %s: IMEI 2 '%s' is already assigned in this transfer.") % (row_num, imei2))

            # Made In country validation
            country = None
            if made_in:
                country = self._find_country(made_in)
                if not country:
                    raise ValidationError(_("Row %s: Country '%s' in 'Made In' not found.") % (row_num, made_in))
            else:
                country = (
                    move.made_in_country_id
                    or move.made_country
                    or (picking.made_country if not is_outgoing else False)
                )

            r["resolved_move"] = move
            r["resolved_country"] = country
            rows_by_move.setdefault(move, []).append(r)

        # 2. Apply to moves and move lines
        MoveLine = self.env["stock.move.line"]

        for move, move_rows in rows_by_move.items():
            product = move.product_id

            if not self.keep_lines:
                # Unlink all move lines for this move if user explicitly unchecked keep_lines
                move.move_line_ids.unlink()
            else:
                # Clean up empty placeholder lines that have no serial/lot and no IMEI
                empty_lines = move.move_line_ids.filtered(
                    lambda l: not l.lot_name and not l.lot_id and not l.imei
                )
                empty_lines.unlink()

            default_vals = {
                "move_id": move.id,
                "picking_id": picking.id,
                "product_id": product.id,
                "product_uom_id": product.uom_id.id,
                "location_id": move.location_id.id,
                "location_dest_id": move.location_dest_id.id,
                "quantity": 1,
            }

            # Putaway strategy resolution for incoming
            if not is_outgoing and move.location_dest_id:
                putaway_loc = move.location_dest_id._get_putaway_strategy(product, 1)
                if putaway_loc and putaway_loc.id:
                    default_vals["location_dest_id"] = putaway_loc.id

            vals_list = []
            for r in move_rows:
                serial = r["serial_number"]
                imei1 = r["imei1"] or False
                imei2 = r["imei2"] or False
                country = r["resolved_country"]

                line_vals = dict(default_vals)
                line_vals["imei"] = imei1
                line_vals["imei2"] = imei2
                if country:
                    line_vals["made_in_country_id"] = country.id
                    line_vals["made_country"] = country.id

                if is_outgoing:
                    # Outgoing delivery: find existing lot in stock
                    lot = self.env["stock.lot"].search([
                        ("product_id", "=", product.id),
                        ("name", "=", serial),
                    ], limit=1)
                    if not lot and imei1:
                        # Try finding by IMEI in stock.quant
                        quant = self.env["stock.quant"].search([
                            ("product_id", "=", product.id),
                            "|", ("imei", "=", imei1), ("imei2", "=", imei1),
                            ("lot_id", "!=", False),
                        ], limit=1)
                        if quant:
                            lot = quant.lot_id

                    if not lot and product.tracking != "none":
                        raise ValidationError(
                            _("Serial number '%s' not found in stock for product '%s'.")
                            % (serial, product.display_name)
                        )
                    line_vals["lot_id"] = lot.id if lot else False
                else:
                    # Incoming receipt: set lot_name
                    line_vals["lot_name"] = serial

                vals_list.append(line_vals)

            # If use_existing_lots is enabled on picking type, resolve lot_ids
            if move.picking_type_id and move.picking_type_id.use_existing_lots and not is_outgoing:
                move._create_lot_ids_from_move_line_vals(vals_list, product.id, move.company_id.id)

            # Create the move lines
            for vals in vals_list:
                MoveLine.create(vals)

            # Update move-level country if not yet set
            if move_rows and move_rows[0]["resolved_country"] and not move.made_in_country_id:
                move.made_in_country_id = move_rows[0]["resolved_country"].id
                move.made_country = move_rows[0]["resolved_country"].id
