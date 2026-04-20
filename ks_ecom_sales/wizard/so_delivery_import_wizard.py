import base64
import logging
import re
from collections import defaultdict
from io import BytesIO

_logger = logging.getLogger(__name__)

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class KsSoDeliveryImportWizard(models.TransientModel):
    _name = "ks.so.delivery.import.wizard"
    _description = "Sale Order Delivery Import Wizard (E-com)"

    file_data = fields.Binary(string="XLSX File", required=True)
    file_name = fields.Char(string="File Name", required=True)

    _ORDER_ID_ALIASES = {"orderid", "order_id", "orderno", "order_no", "ecomorderid"}
    _PRODUCT_CODE_ALIASES = {"asin", "sku", "productcode", "defaultcode", "internalreference"}
    _QTY_DONE_ALIASES = {
        "orderquantity", "qtydone", "quantitydone", "deliveredqty", "qtydelivered",
        "doneqty", "quantity", "qty", "shippedqty",
    }

    @staticmethod
    def _normalize_header(value):
        return re.sub(r"[^a-z0-9]+", "", (value or "").strip().lower())

    @staticmethod
    def _to_string(value):
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _to_float(value, default=0.0):
        if value in (None, ""):
            return default
        try:
            return float(value)
        except Exception as exc:
            raise UserError(_("Invalid number value: %s") % value) from exc

    def _find_header(self, normalized_headers, aliases):
        for alias in aliases:
            if alias in normalized_headers:
                return alias
        return None

    def _extract_sheet_data(self):
        self.ensure_one()
        if load_workbook is None:
            raise UserError(_("Python package 'openpyxl' is required to import XLSX files."))

        if not self.file_name or not self.file_name.lower().endswith(".xlsx"):
            raise UserError(_("Please upload a valid .xlsx file only."))

        try:
            decoded_file = base64.b64decode(self.file_data)
        except Exception as exc:
            raise UserError(_("Unable to decode the uploaded file. Please upload it again.")) from exc

        try:
            workbook = load_workbook(filename=BytesIO(decoded_file), data_only=True)
            sheet = workbook.active
        except Exception as exc:
            raise UserError(_("Unable to read the XLSX file. Please verify file content.")) from exc

        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise UserError(_("The XLSX file does not contain importable rows."))

        headers = [self._to_string(cell) for cell in (rows[0] or [])]
        if not any(headers):
            raise UserError(_("Header row is empty in the uploaded file."))

        normalized_headers = {}
        for idx, header in enumerate(headers):
            normalized = self._normalize_header(header)
            if normalized:
                normalized_headers[normalized] = idx

        order_id_key = self._find_header(normalized_headers, self._ORDER_ID_ALIASES)
        asin_key = self._find_header(normalized_headers, self._PRODUCT_CODE_ALIASES)
        qty_done_key = self._find_header(normalized_headers, self._QTY_DONE_ALIASES)

        if not order_id_key:
            raise UserError(_("Missing required column: Order ID."))
        if not asin_key:
            raise UserError(_("Missing required column: ASIN / Product Code."))
        if not qty_done_key:
            raise UserError(_("Missing required column: Quantity."))

        delivery_data = []
        for line_no, row in enumerate(rows[1:], start=2):
            if not row or all(cell in (None, "") for cell in row):
                continue

            def get_cell(key, _row=row):
                idx = normalized_headers.get(key)
                if idx is None or idx >= len(_row):
                    return None
                return _row[idx]

            order_id = self._to_string(get_cell(order_id_key))
            asin = self._to_string(get_cell(asin_key))
            qty_done = self._to_float(get_cell(qty_done_key), default=0.0)

            if not order_id:
                raise UserError(_("Order ID is missing at row %s.") % line_no)
            if not asin:
                raise UserError(_("ASIN is missing at row %s.") % line_no)
            if qty_done <= 0:
                raise UserError(_("Quantity must be greater than 0 at row %s.") % line_no)

            row_raw = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                row_raw[header] = self._to_string(value)

            delivery_data.append({
                "line_no": line_no,
                "order_id": order_id,
                "asin": asin,
                "qty_done": qty_done,
                "raw": row_raw,
            })

        if not delivery_data:
            raise UserError(_("No valid import lines found in the uploaded file."))

        return delivery_data

    def action_import_delivery_data(self):
        """Update existing sale order delivery (stock picking OUT) quantities. No new deliveries created."""
        self.ensure_one()
        delivery_data = self._extract_sheet_data()

        grouped_data = defaultdict(list)
        for row in delivery_data:
            grouped_data[row["order_id"]].append(row)

        updated_count = 0
        skipped_count = 0
        warnings = []

        for order_id, rows in grouped_data.items():
            # ── Step 1: Find delivery directly by ks_ecom_order_id on stock.picking ──
            picking = self.env["stock.picking"].search([
                ("ks_ecom_order_id", "=", order_id),
                ("picking_type_id.code", "=", "outgoing"),
                ("state", "not in", ("cancel",)),
            ], limit=1)

            # ── Step 2: Fallback — find via Sale Order ──
            if not picking:
                so = self.env["sale.order"].search([
                    ("ks_ecom_order_id", "=", order_id),
                    ("state", "in", ("sale", "done")),
                ], limit=1)

                if not so:
                    warnings.append(_("Order ID '%s' not found. Import the Sale Order file first.") % order_id)
                    skipped_count += len(rows)
                    continue

                pickings = so.picking_ids.filtered(
                    lambda p: p.picking_type_id.code == "outgoing"
                    and p.state not in ("cancel",)
                )
                if not pickings:
                    warnings.append(
                        _("No delivery found for Order ID '%s' (SO: %s). Confirm the sale order first.") % (order_id, so.name)
                    )
                    skipped_count += len(rows)
                    continue

                # Prefer non-done picking; fallback to done
                picking = (
                    pickings.filtered(lambda p: p.state != "done").sorted("create_date", reverse=True)[:1]
                    or pickings.sorted("create_date", reverse=True)[:1]
                )[0]
                # Backfill ks_ecom_order_id on the picking for future imports
                picking.ks_ecom_order_id = order_id

            # Use the picking's company for all stock operations
            company = picking.company_id
            env_company = self.with_company(company).env

            for row in rows:
                asin = row["asin"]
                qty_done = row["qty_done"]

                # Search product within the picking's company context
                product = env_company["product.product"].search(
                    [("default_code", "=", asin)], limit=1
                )
                if not product:
                    # Fallback: search across all companies (shared products)
                    product = self.env["product.product"].sudo().search(
                        [("default_code", "=", asin)], limit=1
                    )
                if not product:
                    warnings.append(
                        _("Product with ASIN '%s' not found (Order ID: %s, Row: %s).") % (asin, order_id, row["line_no"])
                    )
                    skipped_count += 1
                    continue

                # --- Robust move search ---
                # 1. Exact product.product match
                move = picking.move_ids.filtered(
                    lambda m, p=product: m.product_id.id == p.id and m.state != "cancel"
                )
                # 2. Fallback: same product template (handles variant mismatches)
                if not move:
                    tmpl_id = product.product_tmpl_id.id
                    move = picking.move_ids.filtered(
                        lambda m, t=tmpl_id: m.product_id.product_tmpl_id.id == t and m.state != "cancel"
                    )
                # 3. Fallback: match by default_code on the move's product
                if not move:
                    move = picking.move_ids.filtered(
                        lambda m, a=asin: m.product_id.default_code == a and m.state != "cancel"
                    )

                if not move:
                    _logger.info(
                        "No move found for ASIN %s in picking %s — creating new move.",
                        asin, picking.name,
                    )
                    new_move = env_company["stock.move"].create({
                        "name": product.name or asin,
                        "product_id": product.id,
                        "product_uom": product.uom_id.id,
                        "product_uom_qty": qty_done,
                        "quantity": qty_done,
                        "picking_id": picking.id,
                        "location_id": picking.location_id.id,
                        "location_dest_id": picking.location_dest_id.id,
                        "state": "draft",
                    })
                    new_move._action_confirm()
                    move = new_move

                move = move[0]

                # Update or create move line — always in picking's company context
                move_line = move.move_line_ids.filtered(
                    lambda ml, p=product: ml.product_id.id == p.id
                )
                if not move_line:
                    env_company["stock.move.line"].create({
                        "move_id": move.id,
                        "product_id": product.id,
                        "product_uom_id": product.uom_id.id,
                        "location_id": move.location_id.id,
                        "location_dest_id": move.location_dest_id.id,
                        "quantity": qty_done,
                        "picked": True,
                    })
                else:
                    move_line.with_company(company).write({"quantity": qty_done, "picked": True})
                updated_count += 1

        # ── Bus notifications ────────────────────────────────────────────────
        bus = self.env["bus.bus"]
        partner_id = self.env.user.partner_id

        if warnings:
            warning_list = "; ".join(warnings[:10])
            if len(warnings) > 10:
                warning_list += _(" (and %s more)") % (len(warnings) - 10)
            bus._sendone(partner_id, "simple_notification", {
                "type": "danger",
                "title": _("Delivery Import: Records not updated"),
                "message": warning_list,
                "sticky": True,
            })

        if updated_count > 0:
            bus._sendone(partner_id, "simple_notification", {
                "type": "success",
                "title": _("Delivery Import: Success"),
                "message": _("Updated: %s delivery line(s).") % updated_count,
                "sticky": True,
            })
        elif not warnings:
            bus._sendone(partner_id, "simple_notification", {
                "type": "warning",
                "title": _("Delivery Import"),
                "message": _("No lines were updated. Check file format and data."),
                "sticky": True,
            })

        return {"type": "ir.actions.act_window_close"}
