from odoo import models, fields, api
from odoo.exceptions import ValidationError
import base64
import csv
from io import StringIO, BytesIO
from openpyxl import load_workbook
import logging
from datetime import datetime

_logger = logging.getLogger(__name__)


class ActivationStatusWizard(models.TransientModel):
    _name = 'activation.status.wizard'
    _description = 'Activation Status Wizard'

    csv_file = fields.Binary(string='CSV or Excel File', required=True)
    filename = fields.Char(string="Filename")

    def action_process_file(self):
        if not self.csv_file or not self.filename:
            raise ValidationError("Please upload a file.")

        try:
            file_content = base64.b64decode(self.csv_file)
            records = []

            # Determine file type
            if self.filename.lower().endswith('.csv'):
                # CSV processing
                data = StringIO(file_content.decode('utf-8'))
                reader = csv.DictReader(data)
                records = [row for row in reader]

            elif self.filename.lower().endswith(('.xlsx', '.xlsm', '.xls')):
                # Excel processing
                workbook = load_workbook(filename=BytesIO(file_content), data_only=True)
                sheet = workbook.active
                headers = []

                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index == 0:
                        headers = [str(cell).strip() if cell else '' for cell in row]
                    else:
                        row_data = dict(zip(headers, row))
                        records.append(row_data)
            else:
                raise ValidationError("Unsupported file format. Please upload a CSV or Excel file.")

            # ✅ Pass the unified list of dicts to a separate processor
            file_arraay_data = self._process_data(records)
            self._update_status_in_inventory(file_arraay_data)

        except Exception as e:
            raise ValidationError(f"Failed to process file: {str(e)}")

        # self.env.user.notify_success(message="CSV imported successfully! Devices updated.")

        return {'type': 'ir.actions.client', 'tag': 'reload'}

    def _process_data(self, data_list):
        records = []
        for idx, row in enumerate(data_list, start=1):
            try:
                results_raw = row['Result'] 
                if not results_raw:
                    raise ValueError(f"Missing 'Result' field in row {idx}")

                parsed = self._parse_results_string(results_raw)

                imei_raw = parsed.get('IMEI')
                if not imei_raw:
                    raise ValueError(f"Missing 'IMEI' in parsed data at row {idx}")
                imei = imei_raw.strip()

                date_str = parsed.get('Estimated Purchase Date')
                if not date_str:
                    ValueError(f"Missing 'Estimated Purchase Date' in parsed data at row {idx}") 
                activation_date = self._grab_purchase_date(date_str).strip()

                parsed_data_in_dict = {
                    "imei": imei,
                    "activation_date": activation_date
                }
                records.append(parsed_data_in_dict)

            except Exception as e:
                raise ValidationError(f"Failed to process file: {str(e)}")

        return records

    def _parse_results_string(self, results_string):
        data = {}
        if not results_string:
            return data

        try:
            # Split by pipe and loop
            for item in results_string.split('|'):
                if ':' in item:
                    key, value = item.split(':', 1)
                    key = key.strip()
                    value = value.strip()
                    data[key] = value
        except Exception as e:
            raise ValidationError(f"Failed to process file: {str(e)}")

        return data
    
    def _grab_purchase_date(self, date_containing_string):
        if not date_containing_string:
            return ''
        try:
            splits = date_containing_string.split('r')
            return splits[0]
        except Exception as e:
            raise ValidationError(f"Error grabbing purchase date from: {date_containing_string}")

    # Custom override to test


    def _update_status_in_inventory(self, file_arraay_data):

        for idx, item in enumerate(file_arraay_data, start=1):
            imei = item.get('imei')
            activation_date_as_string = item.get('activation_date')
            activation_date = datetime.strptime(activation_date_as_string, '%d %b %Y').date()


            quants = self.env['stock.quant'].search([], order='create_date desc')
            # quants = self.env['stock.quant'].search([])
            for quant in quants:
                if quant.imei == imei or quant.imei2 == imei:
                    quant.write({'activation_status': True, 'activation_date': activation_date})
                    break




        # for idx, item in enumerate(file_arraay_data, start=1):
        #     try:
        #         imei = item.get('imei')
        #         activation_date_as_string = item.get('activation_date')

        #         try:
        #             activation_date = datetime.strptime(activation_date_as_string, '%d %b %Y').date()
        #         except ValueError:
        #             raise ValueError(f"Invalid date format at item {idx}: '{activation_date_as_string}' (expected format: 'dd MMM YYYY')")


        #         quant = self.env['stock.quant'].search([('imei', '=', imei)], limit=1)
        #         if not quant:
        #             quant = self.env['stock.quant'].search([('imei2', '=', imei)], limit=1)

        #         if quant:
        #             quant.write({'activation_status': True, 'activation_date': activation_date})
        #         else:
        #             raise ValueError(f"No stock quant found for IMEI: {imei}")
        #     except Exception as e:
        #         raise ValueError(f"Error updating inventory status for item {idx}: {e}")



 